from flask import Flask, render_template, redirect, url_for, flash, request, Response, send_file
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, login_required, logout_user, current_user
from flask_mail import Mail, Message
from werkzeug.security import generate_password_hash, check_password_hash
from datetime import datetime
from functools import wraps
import os
import io
import threading
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)

# ── Core config ────────────────────────────────────────────
app.config["SECRET_KEY"]                     = "shieldscan_secret_2024"
app.config["SQLALCHEMY_DATABASE_URI"]        = "sqlite:///shieldscan.db"
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

# ── Email config ───────────────────────────────────────────
app.config["MAIL_SERVER"]         = "smtp.gmail.com"
app.config["MAIL_PORT"]           = 587
app.config["MAIL_USE_TLS"]        = True
app.config["MAIL_USERNAME"]       = "your_gmail@gmail.com"
app.config["MAIL_PASSWORD"]       = "your_app_password"
app.config["MAIL_DEFAULT_SENDER"] = "your_gmail@gmail.com"
ALERT_EMAIL_RECIPIENT             = "recipient_email@gmail.com"

db            = SQLAlchemy(app)
login_manager = LoginManager(app)
login_manager.login_view = "login"
mail          = Mail(app)

# ─────────────────────────────────────────
#  MODELS
# ─────────────────────────────────────────

class User(UserMixin, db.Model):
    __tablename__ = "users"
    id         = db.Column(db.Integer,     primary_key=True)
    username   = db.Column(db.String(150), nullable=False)
    email      = db.Column(db.String(150), unique=True, nullable=False)
    role       = db.Column(db.String(100), nullable=False, default="Staff")
    password   = db.Column(db.String(256), nullable=False)
    is_admin   = db.Column(db.Boolean,     default=False)
    created_at = db.Column(db.DateTime,    default=datetime.utcnow)
    is_active  = db.Column(db.Boolean,     default=True)
    logins     = db.relationship("LoginRecord", backref="user",
                                 lazy=True,
                                 foreign_keys="LoginRecord.user_id")

class LoginRecord(db.Model):
    __tablename__ = "login_records"
    id         = db.Column(db.Integer, primary_key=True)
    user_id    = db.Column(db.Integer, db.ForeignKey("users.id"),
                           nullable=False)
    login_time = db.Column(db.DateTime, default=datetime.utcnow)
    ip_address = db.Column(db.String(50))
    status     = db.Column(db.String(20), default="success")

class VisitorLog(db.Model):
    __tablename__ = "visitor_logs"
    id           = db.Column(db.Integer, primary_key=True)
    detected_by  = db.Column(db.Integer, db.ForeignKey("users.id"),
                             nullable=False)
    visitor_name = db.Column(db.String(150), default="Unknown")
    mask_status  = db.Column(db.String(20),  nullable=False)
    confidence   = db.Column(db.Float,       nullable=False)
    location     = db.Column(db.String(150), default="Main Entrance")
    timestamp    = db.Column(db.DateTime,    default=datetime.utcnow)

class ShutdownLog(db.Model):
    __tablename__ = "shutdown_logs"
    id           = db.Column(db.Integer, primary_key=True)
    user_id      = db.Column(db.Integer, db.ForeignKey("users.id"),
                             nullable=False)
    reason       = db.Column(db.String(200),
                             default="No mask detected for 60 seconds")
    nomask_count = db.Column(db.Integer,  default=0)
    duration     = db.Column(db.Integer,  default=60)
    timestamp    = db.Column(db.DateTime, default=datetime.utcnow)
    user         = db.relationship("User", backref="shutdowns",
                                   lazy=True,
                                   foreign_keys=[user_id])

# ─────────────────────────────────────────
#  AUTH LOADER & DECORATORS
# ─────────────────────────────────────────

@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))

def admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not current_user.is_authenticated or not current_user.is_admin:
            flash("Access denied. Admin only.", "danger")
            return redirect(url_for("dashboard"))
        return f(*args, **kwargs)
    return decorated

# ─────────────────────────────────────────
#  EMAIL FUNCTIONS
# ─────────────────────────────────────────

def send_nomask_email(confidence, location, detected_by):
    try:
        msg = Message(
            subject    = "ShieldScan Alert — No Mask Detected!",
            recipients = [ALERT_EMAIL_RECIPIENT]
        )
        msg.html = f"""
        <div style="font-family:Arial,sans-serif; max-width:520px;
                    margin:0 auto; border:1px solid #ddd;
                    border-radius:12px; overflow:hidden;">
          <div style="background:#a32d2d; padding:20px 24px;">
            <h2 style="color:white; margin:0;">
              ShieldScan — No Mask Alert
            </h2>
          </div>
          <div style="padding:24px;">
            <p style="font-size:15px; color:#333;">
              A person without a face mask has been detected.
            </p>
            <table style="width:100%; border-collapse:collapse;
                          font-size:14px; margin-top:16px;">
              <tr style="background:#f8f8f8;">
                <td style="padding:10px 14px; color:#888;
                           border-bottom:1px solid #eee;">
                  Detection time
                </td>
                <td style="padding:10px 14px; color:#333;
                           border-bottom:1px solid #eee;">
                  {datetime.utcnow().strftime("%d %b %Y, %I:%M:%S %p")} UTC
                </td>
              </tr>
              <tr>
                <td style="padding:10px 14px; color:#888;
                           border-bottom:1px solid #eee;">
                  Location
                </td>
                <td style="padding:10px 14px; color:#333;
                           border-bottom:1px solid #eee;">
                  {location}
                </td>
              </tr>
              <tr style="background:#f8f8f8;">
                <td style="padding:10px 14px; color:#888;
                           border-bottom:1px solid #eee;">
                  Confidence
                </td>
                <td style="padding:10px 14px; color:#a32d2d;
                           font-weight:bold;
                           border-bottom:1px solid #eee;">
                  {confidence:.1f}%
                </td>
              </tr>
              <tr>
                <td style="padding:10px 14px; color:#888;">
                  Detected by
                </td>
                <td style="padding:10px 14px; color:#333;">
                  {detected_by}
                </td>
              </tr>
            </table>
            <div style="margin-top:20px; padding:14px;
                        background:#fdecea; border-radius:8px;
                        font-size:13px; color:#a32d2d;">
              Please take immediate action to ensure mask
              compliance at the detected location.
            </div>
          </div>
          <div style="background:#f8f8f8; padding:14px 24px;
                      font-size:12px; color:#aaa; text-align:center;">
            ShieldScan — Real-time Face Mask Detection System
          </div>
        </div>
        """
        mail.send(msg)
        print(f"Alert email sent to {ALERT_EMAIL_RECIPIENT}")
    except Exception as e:
        print(f"Email error: {e}")

def send_shutdown_email(username, role, nomask_count):
    try:
        msg = Message(
            subject    = "ShieldScan — Camera Shutdown Alert!",
            recipients = [ALERT_EMAIL_RECIPIENT]
        )
        msg.html = f"""
        <div style="font-family:Arial,sans-serif; max-width:520px;
                    margin:0 auto; border:1px solid #ddd;
                    border-radius:12px; overflow:hidden;">
          <div style="background:#3C3489; padding:20px 24px;">
            <h2 style="color:white; margin:0;">
              ShieldScan — Camera Shutdown Alert
            </h2>
          </div>
          <div style="padding:24px;">
            <p style="font-size:15px; color:#333;">
              A user camera was automatically shut down due to
              persistent no mask detection.
            </p>
            <table style="width:100%; border-collapse:collapse;
                          font-size:14px; margin-top:16px;">
              <tr style="background:#f8f8f8;">
                <td style="padding:10px 14px; color:#888;
                           border-bottom:1px solid #eee;">
                  User
                </td>
                <td style="padding:10px 14px; color:#333;
                           font-weight:bold;
                           border-bottom:1px solid #eee;">
                  {username}
                </td>
              </tr>
              <tr>
                <td style="padding:10px 14px; color:#888;
                           border-bottom:1px solid #eee;">
                  Role
                </td>
                <td style="padding:10px 14px; color:#333;
                           border-bottom:1px solid #eee;">
                  {role}
                </td>
              </tr>
              <tr style="background:#f8f8f8;">
                <td style="padding:10px 14px; color:#888;
                           border-bottom:1px solid #eee;">
                  No mask detections
                </td>
                <td style="padding:10px 14px; color:#a32d2d;
                           font-weight:bold;
                           border-bottom:1px solid #eee;">
                  {nomask_count} times
                </td>
              </tr>
              <tr>
                <td style="padding:10px 14px; color:#888;">
                  Shutdown time
                </td>
                <td style="padding:10px 14px; color:#333;">
                  {datetime.utcnow().strftime("%d %b %Y, %I:%M:%S %p")} UTC
                </td>
              </tr>
            </table>
            <div style="margin-top:20px; padding:14px;
                        background:#EEEDFE; border-radius:8px;
                        font-size:13px; color:#3C3489;">
              This user repeatedly failed to wear a mask.
              Please follow up with them immediately.
            </div>
          </div>
          <div style="background:#f8f8f8; padding:14px 24px;
                      font-size:12px; color:#aaa; text-align:center;">
            ShieldScan — Real-time Face Mask Detection System
          </div>
        </div>
        """
        mail.send(msg)
        print(f"Shutdown email sent for: {username}")
    except Exception as e:
        print(f"Shutdown email error: {e}")

# ─────────────────────────────────────────
#  EXCEL EXPORT
# ─────────────────────────────────────────

def style_header_row(ws, row, num_cols, bg_color):
    fill   = PatternFill("solid", fgColor=bg_color)
    font   = Font(bold=True, color="FFFFFF", size=11)
    border = Border(bottom=Side(style="thin", color="CCCCCC"))
    for col in range(1, num_cols + 1):
        cell           = ws.cell(row=row, column=col)
        cell.fill      = fill
        cell.font      = font
        cell.alignment = Alignment(horizontal="center",
                                   vertical="center")
        cell.border    = border

def style_data_row(ws, row, num_cols, is_even):
    bg   = "F8F9FA" if is_even else "FFFFFF"
    fill = PatternFill("solid", fgColor=bg)
    font = Font(size=10)
    for col in range(1, num_cols + 1):
        cell           = ws.cell(row=row, column=col)
        cell.fill      = fill
        cell.font      = font
        cell.alignment = Alignment(horizontal="left",
                                   vertical="center")

def auto_fit_columns(ws):
    for col in ws.columns:
        max_len    = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)

def generate_excel():
    wb = openpyxl.Workbook()

    # Sheet 1 — Users
    ws1 = wb.active
    ws1.title = "Registered Users"
    ws1.row_dimensions[1].height = 30
    h1 = ["#", "Username", "Email", "Role",
          "Type", "Registered On", "Status"]
    for col, h in enumerate(h1, 1):
        ws1.cell(row=1, column=col, value=h)
    style_header_row(ws1, 1, len(h1), "0C447C")
    users = User.query.order_by(User.created_at.desc()).all()
    for i, u in enumerate(users, 1):
        row = [
            i, u.username, u.email, u.role,
            "Admin" if u.is_admin else "User",
            u.created_at.strftime("%d %b %Y, %I:%M %p"),
            "Active" if u.is_active else "Inactive"
        ]
        for col, v in enumerate(row, 1):
            ws1.cell(row=i+1, column=col, value=v)
        style_data_row(ws1, i+1, len(h1), i % 2 == 0)
    auto_fit_columns(ws1)

    # Sheet 2 — Login History
    ws2 = wb.create_sheet("Login History")
    ws2.row_dimensions[1].height = 30
    h2 = ["#", "Username", "Email", "Role",
          "Login Time", "IP Address", "Status"]
    for col, h in enumerate(h2, 1):
        ws2.cell(row=1, column=col, value=h)
    style_header_row(ws2, 1, len(h2), "1D9E75")
    logins = LoginRecord.query.order_by(
        LoginRecord.login_time.desc()).all()
    for i, log in enumerate(logins, 1):
        row = [
            i, log.user.username, log.user.email,
            log.user.role,
            log.login_time.strftime("%d %b %Y, %I:%M %p"),
            log.ip_address,
            "Success" if log.status == "success" else "Failed"
        ]
        for col, v in enumerate(row, 1):
            ws2.cell(row=i+1, column=col, value=v)
        style_data_row(ws2, i+1, len(h2), i % 2 == 0)
    auto_fit_columns(ws2)

    # Sheet 3 — Visitor Logs
    ws3 = wb.create_sheet("Visitor Detection Logs")
    ws3.row_dimensions[1].height = 30
    h3 = ["#", "Detected By", "Visitor Name",
          "Mask Status", "Confidence %",
          "Location", "Timestamp"]
    for col, h in enumerate(h3, 1):
        ws3.cell(row=1, column=col, value=h)
    style_header_row(ws3, 1, len(h3), "712B13")
    visitors = VisitorLog.query.order_by(
        VisitorLog.timestamp.desc()).all()
    for i, v in enumerate(visitors, 1):
        try:
            u    = User.query.get(v.detected_by)
            name = u.username if u else "Unknown"
        except Exception:
            name = "Unknown"
        row = [
            i, name, v.visitor_name, v.mask_status,
            round(v.confidence, 1), v.location,
            v.timestamp.strftime("%d %b %Y, %I:%M %p")
        ]
        for col, val in enumerate(row, 1):
            ws3.cell(row=i+1, column=col, value=val)
        style_data_row(ws3, i+1, len(h3), i % 2 == 0)
        cell      = ws3.cell(row=i+1, column=4)
        cell.font = Font(
            color = "1e7e34" if v.mask_status == "Mask" else "a32d2d",
            bold  = True, size=10
        )
    auto_fit_columns(ws3)

    # Sheet 4 — Shutdown Logs
    ws4 = wb.create_sheet("Shutdown Logs")
    ws4.row_dimensions[1].height = 30
    h4 = ["#", "Username", "Email", "Role",
          "No Mask Count", "Reason",
          "Duration (s)", "Shutdown Time"]
    for col, h in enumerate(h4, 1):
        ws4.cell(row=1, column=col, value=h)
    style_header_row(ws4, 1, len(h4), "3C3489")
    shutdowns = ShutdownLog.query.order_by(
        ShutdownLog.timestamp.desc()).all()
    for i, s in enumerate(shutdowns, 1):
        try:
            u    = User.query.get(s.user_id)
            name  = u.username if u else "Unknown"
            email = u.email    if u else "Unknown"
            role  = u.role     if u else "Unknown"
        except Exception:
            name = email = role = "Unknown"
        row = [
            i, name, email, role,
            s.nomask_count, s.reason,
            s.duration,
            s.timestamp.strftime("%d %b %Y, %I:%M %p")
        ]
        for col, val in enumerate(row, 1):
            ws4.cell(row=i+1, column=col, value=val)
        style_data_row(ws4, i+1, len(h4), i % 2 == 0)
        cell      = ws4.cell(row=i+1, column=5)
        cell.font = Font(color="a32d2d", bold=True, size=10)
    auto_fit_columns(ws4)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

# ─────────────────────────────────────────
#  ROUTES — PUBLIC
# ─────────────────────────────────────────

@app.route("/")
def home():
    return redirect(url_for("login"))

@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        username         = request.form.get("username")
        email            = request.form.get("email")
        role             = request.form.get("role")
        password         = request.form.get("password")
        confirm_password = request.form.get("confirm_password")

        if password != confirm_password:
            flash("Passwords do not match.", "danger")
            return redirect(url_for("register"))
        if len(password) < 6:
            flash("Password must be at least 6 characters.", "danger")
            return redirect(url_for("register"))

        existing = User.query.filter_by(email=email).first()
        if existing:
            flash("Email already registered.", "danger")
            return redirect(url_for("login"))

        new_user = User(
            username = username,
            email    = email,
            role     = role,
            password = generate_password_hash(password),
            is_admin = False
        )
        db.session.add(new_user)
        db.session.commit()
        flash("Account created! Please login.", "success")
        return redirect(url_for("login"))

    return render_template("register.html")

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        email    = request.form.get("email")
        password = request.form.get("password")
        user     = User.query.filter_by(email=email).first()

        if user and check_password_hash(user.password, password):
            login_user(user)
            db.session.add(LoginRecord(
                user_id    = user.id,
                ip_address = request.remote_addr,
                status     = "success"
            ))
            db.session.commit()
            flash(f"Welcome back, {user.username}!", "success")
            if user.is_admin:
                return redirect(url_for("admin_dashboard"))
            return redirect(url_for("dashboard"))
        else:
            if user:
                db.session.add(LoginRecord(
                    user_id    = user.id,
                    ip_address = request.remote_addr,
                    status     = "failed"
                ))
                db.session.commit()
            flash("Invalid email or password.", "danger")

    return render_template("login.html")

@app.route("/logout")
@login_required
def logout():
    logout_user()
    flash("Logged out successfully.", "success")
    return redirect(url_for("login"))

# ─────────────────────────────────────────
#  ROUTES — USER
# ─────────────────────────────────────────

@app.route("/dashboard")
@login_required
def dashboard():
    if current_user.is_admin:
        return redirect(url_for("admin_dashboard"))

    my_detections  = VisitorLog.query.filter_by(
                         detected_by=current_user.id).count()
    my_mask_ok     = VisitorLog.query.filter_by(
                         detected_by=current_user.id,
                         mask_status="Mask").count()
    my_no_mask     = VisitorLog.query.filter_by(
                         detected_by=current_user.id,
                         mask_status="No Mask").count()
    my_logins      = LoginRecord.query.filter_by(
                         user_id=current_user.id).count()
    my_shutdowns   = ShutdownLog.query.filter_by(
                         user_id=current_user.id).count()
    recent_my_logs = VisitorLog.query.filter_by(
                         detected_by=current_user.id
                     ).order_by(
                         VisitorLog.timestamp.desc()
                     ).limit(5).all()

    return render_template("user_dashboard.html",
        username       = current_user.username,
        role           = current_user.role,
        my_detections  = my_detections,
        my_mask_ok     = my_mask_ok,
        my_no_mask     = my_no_mask,
        my_logins      = my_logins,
        my_shutdowns   = my_shutdowns,
        recent_my_logs = recent_my_logs
    )

@app.route("/my_records")
@login_required
def my_records():
    if current_user.is_admin:
        return redirect(url_for("records"))

    my_logins    = LoginRecord.query.filter_by(
                       user_id=current_user.id
                   ).order_by(
                       LoginRecord.login_time.desc()
                   ).all()
    my_visitors  = VisitorLog.query.filter_by(
                       detected_by=current_user.id
                   ).order_by(
                       VisitorLog.timestamp.desc()
                   ).all()
    my_shutdowns = ShutdownLog.query.filter_by(
                       user_id=current_user.id
                   ).order_by(
                       ShutdownLog.timestamp.desc()
                   ).all()

    return render_template("user_records.html",
        username     = current_user.username,
        role         = current_user.role,
        my_logins    = my_logins,
        my_visitors  = my_visitors,
        my_shutdowns = my_shutdowns
    )

# ─────────────────────────────────────────
#  ROUTES — ADMIN ONLY
# ─────────────────────────────────────────

@app.route("/admin")
@login_required
@admin_required
def admin_dashboard():
    total_users     = User.query.count()
    total_logins    = LoginRecord.query.filter_by(
                          status="success").count()
    total_visitors  = VisitorLog.query.count()
    mask_ok         = VisitorLog.query.filter_by(
                          mask_status="Mask").count()
    no_mask         = VisitorLog.query.filter_by(
                          mask_status="No Mask").count()
    total_shutdowns = ShutdownLog.query.count()
    recent_logins   = LoginRecord.query.order_by(
                          LoginRecord.login_time.desc()
                      ).limit(8).all()
    recent_visitors = VisitorLog.query.order_by(
                          VisitorLog.timestamp.desc()
                      ).limit(8).all()
    shutdown_logs   = ShutdownLog.query.order_by(
                          ShutdownLog.timestamp.desc()
                      ).all()

    return render_template("admin_dashboard.html",
        username        = current_user.username,
        total_users     = total_users,
        total_logins    = total_logins,
        total_visitors  = total_visitors,
        mask_ok         = mask_ok,
        no_mask         = no_mask,
        total_shutdowns = total_shutdowns,
        recent_logins   = recent_logins,
        recent_visitors = recent_visitors,
        shutdown_logs   = shutdown_logs
    )

@app.route("/records")
@login_required
@admin_required
def records():
    users         = User.query.order_by(User.created_at.desc()).all()
    logins        = LoginRecord.query.order_by(
                        LoginRecord.login_time.desc()).all()
    visitors      = VisitorLog.query.order_by(
                        VisitorLog.timestamp.desc()).all()
    shutdown_logs = ShutdownLog.query.order_by(
                        ShutdownLog.timestamp.desc()).all()
    return render_template("records.html",
        users         = users,
        logins        = logins,
        visitors      = visitors,
        shutdown_logs = shutdown_logs
    )

@app.route("/manage_users")
@login_required
@admin_required
def manage_users():
    users = User.query.order_by(User.created_at.desc()).all()
    return render_template("manage_users.html", users=users)

@app.route("/toggle_user/<int:user_id>")
@login_required
@admin_required
def toggle_user(user_id):
    user = User.query.get_or_404(user_id)
    if user.id == current_user.id:
        flash("You cannot deactivate your own account.", "danger")
        return redirect(url_for("manage_users"))
    user.is_active = not user.is_active
    db.session.commit()
    status = "activated" if user.is_active else "deactivated"
    flash(f"{user.username} has been {status}.", "success")
    return redirect(url_for("manage_users"))

@app.route("/toggle_admin/<int:user_id>")
@login_required
@admin_required
def toggle_admin(user_id):
    user = User.query.get_or_404(user_id)
    if user.id == current_user.id:
        flash("You cannot change your own admin status.", "danger")
        return redirect(url_for("manage_users"))
    user.is_admin = not user.is_admin
    db.session.commit()
    status = "promoted to Admin" if user.is_admin else "changed to User"
    flash(f"{user.username} has been {status}.", "success")
    return redirect(url_for("manage_users"))

@app.route("/delete_user/<int:user_id>")
@login_required
@admin_required
def delete_user(user_id):
    user = User.query.get_or_404(user_id)
    if user.id == current_user.id:
        flash("You cannot delete your own account.", "danger")
        return redirect(url_for("manage_users"))
    LoginRecord.query.filter_by(user_id=user.id).delete()
    VisitorLog.query.filter_by(detected_by=user.id).delete()
    ShutdownLog.query.filter_by(user_id=user.id).delete()
    db.session.delete(user)
    db.session.commit()
    flash(f"{user.username} has been deleted.", "success")
    return redirect(url_for("manage_users"))

@app.route("/export_excel")
@login_required
@admin_required
def export_excel():
    try:
        output   = generate_excel()
        filename = f"ShieldScan_Records_{datetime.now().strftime('%d%b%Y_%H%M')}.xlsx"
        return send_file(
            output,
            mimetype      = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment = True,
            download_name = filename
        )
    except Exception as e:
        flash(f"Export failed: {e}", "danger")
        return redirect(url_for("records"))

@app.route("/email_settings", methods=["GET", "POST"])
@login_required
@admin_required
def email_settings():
    global ALERT_EMAIL_RECIPIENT
    if request.method == "POST":
        ALERT_EMAIL_RECIPIENT = request.form.get(
            "recipient_email", ALERT_EMAIL_RECIPIENT
        )
        flash(f"Alert email updated to {ALERT_EMAIL_RECIPIENT}", "success")
        return redirect(url_for("email_settings"))
    return render_template("email_settings.html",
        recipient=ALERT_EMAIL_RECIPIENT
    )

@app.route("/test_email")
@login_required
@admin_required
def test_email():
    try:
        send_nomask_email(98.5, "Main Entrance", current_user.username)
        flash("Test email sent successfully!", "success")
    except Exception as e:
        flash(f"Email failed: {e}", "danger")
    return redirect(url_for("email_settings"))

# ─────────────────────────────────────────
#  ROUTES — SHARED
# ─────────────────────────────────────────

@app.route("/detect")
@login_required
def detect():
    return render_template("detect.html")

@app.route("/video_feed")
@login_required
def video_feed():
    from detect import generate_frames
    return Response(
        generate_frames(user_id=current_user.id),
        mimetype="multipart/x-mixed-replace; boundary=frame"
    )

@app.route("/log_detection", methods=["POST"])
@login_required
def log_detection():
    mask_status = request.form.get("mask_status", "Unknown")
    confidence  = float(request.form.get("confidence", 0.0))
    location    = request.form.get("location", "Main Entrance")

    db.session.add(VisitorLog(
        detected_by  = current_user.id,
        mask_status  = mask_status,
        confidence   = confidence,
        location     = location
    ))
    db.session.commit()

    if mask_status == "No Mask":
        threading.Thread(
            target = send_nomask_email,
            args   = (confidence, location, current_user.username),
            daemon = True
        ).start()

    return {"status": "logged"}, 200

@app.route("/log_shutdown", methods=["POST"])
@login_required
def log_shutdown():
    nomask_count = int(request.form.get("nomask_count", 0))
    duration     = int(request.form.get("duration", 60))

    db.session.add(ShutdownLog(
        user_id      = current_user.id,
        nomask_count = nomask_count,
        duration     = duration
    ))
    db.session.commit()

    threading.Thread(
        target = send_shutdown_email,
        args   = (current_user.username,
                  current_user.role,
                  nomask_count),
        daemon = True
    ).start()

    return {"status": "shutdown_logged"}, 200

@app.route("/session_stats")
@login_required
def session_stats():
    from detect import get_session_stats
    return get_session_stats()

@app.route("/reset_stats")
@login_required
def reset_stats():
    from detect import reset_session_stats
    reset_session_stats()
    return {"status": "reset"}, 200

# ─────────────────────────────────────────
#  CREATE DEFAULT ADMIN
# ─────────────────────────────────────────

def create_default_admin():
    admin = User.query.filter_by(
        email="admin@shieldscan.com"
    ).first()
    if not admin:
        admin = User(
            username = "Admin",
            email    = "admin@shieldscan.com",
            role     = "Administrator",
            password = generate_password_hash("admin123"),
            is_admin = True
        )
        db.session.add(admin)
        db.session.commit()
        print("Default admin created!")
        print("Email    : admin@shieldscan.com")
        print("Password : admin123")

# ─────────────────────────────────────────
#  RUN
# ─────────────────────────────────────────

if __name__ == "__main__":
    with app.app_context():
        db.create_all()
        create_default_admin()
        print("Database ready!")
    app.run(debug=True, threaded=True)